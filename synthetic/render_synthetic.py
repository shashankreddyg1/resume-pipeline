#!/usr/bin/env python3
"""Step 4 — dedicated renderer for synthetic whole-resume JSON.

Separate from the production renderer (render.py); it only REUSES that
module's Jake's-style formatting components (styles, section(), bullets(),
two_col()). It consumes the validated synthetic schema from validate.py.

The SYNTHETIC identification is enforced HERE, in code — not by the LLM:
  - a banner block directly under the header of every rendered resume
  - a footer line on every PDF page
  - "(synthetic — unverified)" appended to every certification
  - document metadata titles prefixed with SYNTHETIC TEST RESUME
No raw JSON, generation metadata, or internal instructions are rendered.
"""
import sys
from pathlib import Path
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import render as R  # production formatting components (styles/helpers only)
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER

BANNER_TEXT = "SYNTHETIC TEST RESUME — NOT FOR REAL-WORLD APPLICATIONS"
FOOTER_TEXT = "SYNTHETIC TEST DATA — NOT A REAL CANDIDATE RESUME"
STRIP_BG = HexColor("#111827")

s_banner = R.st("sb", fontName=R.SERIF_B, fontSize=10, leading=13,
                alignment=TA_CENTER, textColor=HexColor("#ffffff"))


def _esc(s):
    return escape(str(s))


def _banner_flowable():
    """Compact header strip: one line of real text (ATS-extractable),
    white on dark slate — professional, unmissable, not a warning box."""
    t = Table([[Paragraph(BANNER_TEXT, s_banner)]], colWidths=[R.PAGE_W])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), STRIP_BG),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Times-Roman", 8)
    canvas.setFillColor(HexColor("#6b7280"))
    canvas.drawCentredString(LETTER[0] / 2, 0.32 * inch, FOOTER_TEXT)
    canvas.restoreState()


def _contact_lines(c):
    bits = [c.get("location"), c.get("phone"), c.get("email")]
    line1 = " &nbsp;|&nbsp; ".join(_esc(b) for b in bits if b)
    links = []
    for label, key in (("GitHub", "github"), ("Portfolio", "portfolio"),
                       ("LinkedIn", "linkedin")):
        if c.get(key):
            links.append(f'<a href="{_esc(c[key])}" color="#1a56db"><u>{label}</u></a>')
    return line1, " &nbsp;|&nbsp; ".join(links)


def render_pdf(data, out_path):
    story = [Paragraph(_esc(data["name"]), R.s_name)]
    if data.get("title"):
        story.append(Paragraph(_esc(data["title"]), R.s_title))
    line1, line2 = _contact_lines(data.get("contact", {}))
    contact_html = line1 + (("<br/>" + line2) if line2 else "")
    story.append(Paragraph(contact_html, R.s_contact))
    story.append(_banner_flowable())
    story.append(Spacer(1, 6))

    story += R.section("Professional Summary")
    story.append(Paragraph(_esc(data["professional_summary"]), R.s_sum))

    story += R.section("Technical Skills")
    skills = data["technical_skills"]
    for i in range(0, len(skills), 6):
        story.append(Paragraph(", ".join(_esc(s) for s in skills[i:i + 6]),
                               R.s_skill))

    story += R.section("Professional Experience")
    for e in data["experience"]:
        story.append(R.two_col(_esc(e["company"]), _esc(e.get("dates", "")),
                               R.s_co, R.s_dates))
        story.append(R.two_col(_esc(e.get("title") or e.get("role", "")),
                               _esc(e.get("location", "")), R.s_role, R.s_loc))
        story += R.bullets([_esc(b) for b in e["bullets"]])

    if data.get("projects"):
        story += R.section("Selected Projects")
        for p in data["projects"]:
            if p.get("name"):
                story.append(Paragraph(f"<b>{_esc(p['name'])}</b>", R.s_proj))
            links = [(l[0], l[1]) for l in p.get("links", [])
                     if isinstance(l, (list, tuple)) and len(l) == 2]
            if links:
                story.append(R.link_line([(str(a), str(b)) for a, b in links]))
            if p.get("objective"):
                story.append(Paragraph(_esc(p["objective"]), R.s_sum))
            if p.get("architecture"):
                story.append(Paragraph(
                    f"<b>Architecture:</b> {_esc(p['architecture'])}", R.s_sum))
            story += R.bullets([_esc(b) for b in p.get("bullets", [])])
            if p.get("outcome"):
                story.append(Paragraph(
                    f"<b>Outcome:</b> {_esc(p['outcome'])}", R.s_sum))
            if p.get("stack"):
                story.append(Paragraph(
                    "<b>Tech stack:</b> " +
                    ", ".join(_esc(s) for s in p["stack"]), R.s_link))

    if data.get("certifications"):
        story += R.section("Certifications")
        names = [c if isinstance(c, str) else c.get("name", "")
                 for c in data["certifications"]]
        story += R.bullets([f"{_esc(n)} (synthetic — unverified)"
                             for n in names if n])

    story += R.section("Education")
    for ed in data["education"]:
        txt = f"<b>{_esc(ed['school'])}</b> — {_esc(ed['degree'])}" + \
              (f", {_esc(ed['dates'])}" if ed.get("dates") else "")
        story.append(Paragraph(txt, R.s_edu))

    doc = SimpleDocTemplate(out_path, pagesize=LETTER,
                            leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch,
                            title=f"SYNTHETIC TEST RESUME — {data['name']}",
                            author="Synthetic test pipeline")
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return out_path


def render_docx(data, out_path):
    """DOCX export (python-docx). Same content + banner discipline."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.6)
        sec.top_margin = sec.bottom_margin = Inches(0.5)
        fp = sec.footer.paragraphs[0]
        fp.text = FOOTER_TEXT
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)

    def para(text="", bold=False, size=10.5, center=False, italic=False,
             color=None, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11.5, space_after=3)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "111827")
        pbdr.append(bottom); pPr.append(pbdr)

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)

    para(data["name"], bold=True, size=18, center=True, space_after=1)
    if data.get("title"):
        para(data["title"], bold=True, size=11.5, center=True, space_after=1)
    c = data.get("contact", {})
    para(" | ".join(str(x) for x in
                    (c.get("location"), c.get("phone"), c.get("email")) if x),
         size=10, center=True, space_after=6)

    # header strip: compact one-cell table, dark slate fill, white bold text
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "111827")
    cell._tc.get_or_add_tcPr().append(shd)
    bp = cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run(BANNER_TEXT)
    br.bold = True; br.font.size = Pt(10)
    br.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    para("", size=4, space_after=2)

    heading("Professional Summary")
    para(data["professional_summary"])
    heading("Technical Skills")
    para(", ".join(data["technical_skills"]))
    heading("Professional Experience")
    for e in data["experience"]:
        para(f'{e["company"]}' + (f'  —  {e["dates"]}' if e.get("dates") else ""),
             bold=True)
        para(e.get("title") or e.get("role", ""), italic=True, space_after=1)
        for b in e["bullets"]:
            bullet(b)
    if data.get("projects"):
        heading("Selected Projects")
        for p in data["projects"]:
            if p.get("name"):
                para(p["name"], bold=True, space_after=1)
            if p.get("objective"):
                para(p["objective"])
            if p.get("architecture"):
                para("Architecture: " + p["architecture"])
            for b in p.get("bullets", []):
                bullet(b)
            if p.get("outcome"):
                para("Outcome: " + p["outcome"])
            if p.get("stack"):
                para("Tech stack: " + ", ".join(p["stack"]), size=9.5)
    if data.get("certifications"):
        heading("Certifications")
        for cert in data["certifications"]:
            n = cert if isinstance(cert, str) else cert.get("name", "")
            if n:
                bullet(f"{n} (synthetic — unverified)")
    heading("Education")
    for ed in data["education"]:
        para(f'{ed["school"]} — {ed["degree"]}' +
             (f', {ed["dates"]}' if ed.get("dates") else ""))

    core = doc.core_properties
    core.title = f"SYNTHETIC TEST RESUME — {data['name']}"
    core.author = "Synthetic test pipeline"
    doc.save(out_path)
    return out_path
